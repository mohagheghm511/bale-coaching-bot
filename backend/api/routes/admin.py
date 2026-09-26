from sqlalchemy.orm import Session
from sqlalchemy import func
from models.user import User, UserStatus
from models.product import Reservation, ReservationStatus, TimeSlot, Product, ProductType, CoachingPackage
from models.payment import Payment, PaymentStatus, BroadcastMessage
from models.test import Test, Question, Option, ScoreRange, TestResult
from models.static_content import StaticContent, ContentType, MediaType
from models.discount import DiscountCode
from models.installment import InstallmentPlan, Installment
from bale import Bot
from bot.sender import send_message, inline_keyboard, send_photo
from bot.dispatcher import get_state, set_state, clear_state
from datetime import datetime, timedelta

bot_settings = {
    "welcome_message": "????! ?? ???? ?????? ??? ?????",
    "card_number": "",
    "card_owner": "",
}

async def show_dashboard(bot, user, db):
    total_users = db.query(func.count(User.id)).scalar()
    pending_payments = db.query(func.count(Payment.id)).filter(Payment.status == PaymentStatus.uploaded).scalar()
    active_reservations = db.query(func.count(Reservation.id)).filter(Reservation.status == ReservationStatus.confirmed).scalar()
    total_income = db.query(func.sum(Payment.amount)).filter(Payment.status == PaymentStatus.approved).scalar() or 0
    total_tests = db.query(func.count(TestResult.id)).scalar()
    from models.review import Review
    pending_reviews = db.query(func.count(Review.id)).filter(Review.is_approved == False).scalar()
    msg = (
        "?? ??????? ??????\n?????????????\n"
        f"?? ?? ???????: {total_users:,}\n"
        f"?? ??????? ????: {active_reservations:,}\n"
        f"?? ????? ??: {total_income:,} ?????\n"
        f"?? ??? ????? ???: {total_tests:,}\n"
        f"?? ??? ?? ??????: {pending_payments}\n"
        f"? ??? ?? ??????: {pending_reviews}\n?????????????"
    )
    await send_message(bot, user.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ??????? ?? ??????", "callback_data": "admin_pending_payments"},
         {"text": "?? ?????? ??????", "callback_data": "admin_reservations"}],
        [{"text": "?? ?????? ?????", "callback_data": "admin_search_user"},
         {"text": "?? ???? ?????", "callback_data": "admin_broadcast"}],
        [{"text": "?? ?????? ??????", "callback_data": "admin_tests"},
         {"text": "?? ?????? ???????", "callback_data": "admin_products"}],
        [{"text": "? ?????? ???? ????", "callback_data": "admin_add_slot"},
         {"text": "?? ?????? ???????", "callback_data": "admin_packages"}],
        [{"text": "? ????? ?????", "callback_data": "admin_reviews"},
         {"text": "?? ????? ????", "callback_data": "admin_report"}],
        [{"text": "?? ?????? ?????", "callback_data": "admin_static_content"},
         {"text": "?? ???? ???????", "callback_data": "admin_export_users"}],
        [{"text": "?? ????? ?????", "callback_data": "admin_discounts"},
         {"text": "?? ???????", "callback_data": "admin_settings"}],
    ]))

async def show_pending_payments(bot, admin, db):
    payments = db.query(Payment).filter(Payment.status == PaymentStatus.uploaded).order_by(Payment.created_at).all()
    if not payments:
        await send_message(bot, admin.bale_id, "? ??? ???? ?? ?????? ????? ????.")
        return
    await send_message(bot, admin.bale_id, f"?? {len(payments)} ??? ?? ?????? ?????:")
    for p in payments[:5]:
        elapsed = int((datetime.now() - p.created_at.replace(tzinfo=None)).total_seconds() / 60)
        msg = (
            "?????????????\n"
            f"?? {p.user.full_name} | ?? {p.user.phone}\n"
            f"?? {p.amount:,} ????? | ?? {p.tracking_code}\n"
            f"? {elapsed} ????? ???"
        )
        await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
            [{"text": "? ?????", "callback_data": f"admin_approve_pay_{p.id}"},
             {"text": "? ??", "callback_data": f"admin_reject_pay_{p.id}"}],
            [{"text": "?? ?????? ???", "callback_data": f"admin_view_receipt_{p.id}"}],
        ]))

async def show_reservations(bot, admin, db):
    reservations = db.query(Reservation).filter(
        Reservation.status == ReservationStatus.confirmed
    ).order_by(Reservation.created_at.desc()).limit(8).all()
    if not reservations:
        await send_message(bot, admin.bale_id, "???? ????? ???? ?????.")
        return
    await send_message(bot, admin.bale_id, f"?? {len(reservations)} ???? ????:")
    for res in reservations:
        slot_date = res.time_slot.date.strftime("%Y/%m/%d — %H:%M") if res.time_slot else "??????"
        msg = f"?? {res.user.full_name}\n?? {res.session_type}\n?? {slot_date}"
        await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
            [{"text": "?? ???????", "callback_data": f"admin_reschedule_{res.id}"},
             {"text": "? ???", "callback_data": f"admin_cancel_res_{res.id}"}],
        ]))

async def prompt_search_user(bot, admin):
    set_state(admin.bale_id, {"step": "admin_search_user"})
    await send_message(bot, admin.bale_id, "?? ??? ?? ????? ????? ?? ?????:")

async def search_user(bot, admin, query, db):
    users = db.query(User).filter(
        (User.full_name.ilike(f"%{query}%")) |
        (User.phone.ilike(f"%{query}%")) |
        (User.username.ilike(f"%{query}%"))
    ).limit(5).all()
    clear_state(admin.bale_id)
    if not users:
        await send_message(bot, admin.bale_id, "?????? ???? ???.")
        return
    for u in users:
        msg = f"?? {u.full_name}\n?? {u.phone or '??? ????'}\n?? {len(u.test_results)} ??? | ?? {len(u.reservations)} ????"
        await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
            [{"text": "?? ??????? ????", "callback_data": f"admin_user_{u.id}"},
             {"text": "?? ????? ????", "callback_data": f"admin_msg_user_{u.id}"}],
            [{"text": "?? ????", "callback_data": f"admin_block_{u.id}"},
             {"text": "? ??????", "callback_data": f"admin_unblock_{u.id}"}],
        ]))

async def show_user_profile(bot, admin, target_user_id, db):
    u = db.query(User).filter(User.id == target_user_id).first()
    if not u:
        return
    total_paid = sum(p.amount for p in u.payments if p.status == PaymentStatus.approved)
    last_test = u.test_results[-1].test.title if u.test_results and u.test_results[-1].test else "?????"
    msg = (
        "?? ??????? ????\n?????????????\n"
        f"???: {u.full_name}\n??????: {u.phone or '??? ????'}\n"
        f"?????: {'? ????' if u.status == UserStatus.active else '?? ????'}\n"
        "?????????????\n"
        f"?? ??????: {len(u.test_results)}\n"
        f"?? ??????: {len(u.reservations)}\n"
        f"?? ????? ????: {total_paid:,} ?????\n"
        f"?? ?????: {u.created_at.strftime('%Y/%m/%d')}"
    )
    await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ????? ???? ??????", "callback_data": f"admin_msg_user_{u.id}"}],
        [{"text": "?? ????", "callback_data": f"admin_block_{u.id}"},
         {"text": "? ??????", "callback_data": f"admin_unblock_{u.id}"}],
    ]))

async def prompt_direct_message(bot, admin, target_user_id):
    set_state(admin.bale_id, {"step": "admin_direct_msg", "target_user_id": target_user_id})
    await send_message(bot, admin.bale_id, "?? ???? ?? ?????:")

async def send_direct_message(bot, admin, text, db):
    state = get_state(admin.bale_id)
    target_id = state.get("target_user_id")
    clear_state(admin.bale_id)
    target = db.query(User).filter(User.id == target_id).first()
    if not target:
        return
    await send_message(bot, target.bale_id, f"?? ???? ?? ??????:\n\n{text}")
    await send_message(bot, admin.bale_id, f"? ???? ?? {target.full_name} ????? ??.")

async def prompt_broadcast(bot, admin):
    set_state(admin.bale_id, {"step": "admin_broadcast"})
    await send_message(bot, admin.bale_id, "?? ??? ???? ????? ?? ?????:")

async def send_broadcast(bot, admin, text, db):
    clear_state(admin.bale_id)
    users = db.query(User).filter(User.status == UserStatus.active).all()
    await send_message(bot, admin.bale_id, f"? ?? ??? ????? ?? {len(users)} ?????...")
    sent = failed = 0
    for u in users:
        try:
            await send_message(bot, u.bale_id, text)
            sent += 1
        except:
            failed += 1
    log = BroadcastMessage(title="???? ?????", text=text, sent_count=sent, failed_count=failed)
    db.add(log)
    db.commit()
    await send_message(bot, admin.bale_id, f"? ????? ???? ??\n?? ????: {sent} | ? ??????: {failed}")

async def export_users_excel(bot, admin, db):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        import io, tempfile, os
        from datetime import datetime as dt

        users = db.query(User).order_by(User.created_at).all()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "users"
        ws.sheet_view.rightToLeft = True

        header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
        header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        center = Alignment(horizontal="center", vertical="center")
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        even_fill = PatternFill(start_color="F0F8FF", end_color="F0F8FF", fill_type="solid")

        headers = ["row", "name", "phone", "username", "date"]
        widths = [8, 30, 18, 20, 20]
        for col, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = w
        ws.row_dimensions[1].height = 25

        for i, u in enumerate(users, 1):
            row = i + 1
            date_str = str(u.created_at)[:10] if u.created_at else "-"
            data = [i, u.full_name or "-", u.phone or "-",
                    f"@{u.username}" if u.username else "-", date_str]
            for col, val in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = Font(name="Arial", size=10)
                cell.alignment = center
                cell.border = border
                if i % 2 == 0:
                    cell.fill = even_fill
            ws.row_dimensions[row].height = 20

        ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(buf.read())
            tmp_path = tmp.name

        await send_message(bot, admin.bale_id, f"?? ??? ????? ???? ???? {len(users)} ?????...")
        from bale import InputFile
        with open(tmp_path, "rb") as f:
            file_bytes = f.read()
        os.unlink(tmp_path)
        await bot.send_document(
            admin.bale_id,
            InputFile(file_bytes, file_name="users.xlsx"),
            caption=f"users: {len(users)}\ndate: {dt.now().strftime('%Y-%m-%d')}"
        )
    except ImportError:
        await send_message(bot, admin.bale_id, "openpyxl not installed")
    except Exception as e:
        await send_message(bot, admin.bale_id, f"error: {e}")

TYPE_LABELS = {
    "vip": "?? ???? VIP", "individual": "standard",
    "individual2": "intro", "group": "group", "all": "all",
}

async def prompt_add_slot(bot, admin):
    set_state(admin.bale_id, {"step": "admin_slot_type"})
    await send_message(bot, admin.bale_id, "? ??? ?????? ?? ?????? ??:", reply_markup=inline_keyboard([
        [{"text": "?? ???? VIP", "callback_data": "admin_slottype_vip"}],
        [{"text": "?? ???? ?????????", "callback_data": "admin_slottype_individual"}],
        [{"text": "?? ???? ??????", "callback_data": "admin_slottype_individual2"}],
        [{"text": "?? ???? ?????", "callback_data": "admin_slottype_group"}],
        [{"text": "?? ??? ?????", "callback_data": "admin_slottype_all"}],
    ]))

async def prompt_add_slot_datetime(bot, admin, slot_type):
    set_state(admin.bale_id, {"step": "admin_slot_pkg_title", "slot_type": slot_type})
    await send_message(bot, admin.bale_id, f"???: {slot_type}\n????? ???? ?? ?????:")

async def slot_set_pkg_title(bot, admin, text):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_slot_pkg_sessions", "pkg_title": text})
    await send_message(bot, admin.bale_id, "????? ?????:\n(????: 4)")

async def slot_set_pkg_sessions(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        sessions = int(text.strip())
        set_state(admin.bale_id, {**state, "step": "admin_slot_pkg_duration", "pkg_sessions": sessions})
        await send_message(bot, admin.bale_id, "??? ?? ???? (?????):\n(????: 60)")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ???? ??.")

async def slot_set_pkg_duration(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        duration = int(text.strip())
        set_state(admin.bale_id, {**state, "step": "admin_slot_pkg_price", "pkg_duration": duration})
        await send_message(bot, admin.bale_id, "???? ???? (?????):\n(????: 4800000)")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ???? ??.")

async def slot_set_pkg_price(bot, admin, text, db):
    state = get_state(admin.bale_id)
    try:
        price = int(text.replace(",", "").strip())
        set_state(admin.bale_id, {**state, "step": "admin_slot_pkg_capacity", "pkg_price": price})
        await send_message(bot, admin.bale_id, "????? ???? (???):\n(????: 1)")
    except:
        await send_message(bot, admin.bale_id, "???? ?????? ???.")

async def slot_set_pkg_capacity(bot, admin, text, db):
    state = get_state(admin.bale_id)
    try:
        capacity = int(text.strip())
        price = state["pkg_price"]
        count = db.query(CoachingPackage).count()
        pkg = CoachingPackage(
            title=state["pkg_title"], session_type=state.get("slot_type", "individual"),
            session_count=state["pkg_sessions"], session_duration=state["pkg_duration"],
            price=price, capacity=capacity, is_active=True, order=count,
        )
        db.add(pkg)
        db.commit()
        db.refresh(pkg)
        if price > 0:
            set_state(admin.bale_id, {**state, "pkg_id_new": pkg.id, "capacity": capacity, "pkg_id": pkg.id})
            await prompt_installment_question(bot, admin, "package", pkg.id, price)
        else:
            total_slots = state["pkg_sessions"] * capacity
            set_state(admin.bale_id, {
                **state, "step": "admin_slot_datetime", "pkg_id": pkg.id,
                "current_session": 1, "current_person": 1,
                "collected_slots": [], "capacity": capacity,
                "total_slots": total_slots,
            })
            await send_message(bot, admin.bale_id,
                f"???? ????? ??!\n"
                f"??? 1 ?? {capacity}\n"
                f"???? 1 ?? {state['pkg_sessions']}\n"
                f"????: 1405/05/20 16:00")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ???? ??.")

async def add_time_slot(bot, admin, text, db):
    import jdatetime
    state = get_state(admin.bale_id)
    pkg_id = state.get("pkg_id")
    total_sessions = state.get("pkg_sessions", 1)
    current_session = state.get("current_session", 1)
    collected_slots = state.get("collected_slots", [])
    try:
        jalali_dt = jdatetime.datetime.strptime(text.strip(), "%Y/%m/%d %H:%M")
        gregorian_dt = jalali_dt.togregorian()
        slot = TimeSlot(date=gregorian_dt, is_available=True, package_id=pkg_id)
        db.add(slot)
        db.commit()
        db.refresh(slot)
        collected_slots.append({"session": current_session, "date": text, "slot_id": slot.id})
        if current_session < total_sessions:
            next_session = current_session + 1
            set_state(admin.bale_id, {**state, "current_session": next_session, "collected_slots": collected_slots})
            await send_message(bot, admin.bale_id,
                f"???? {current_session}: {text} ??? ??.\n???? {next_session} ?? {total_sessions}\n????: 1405/05/20 16:00")
        else:
            clear_state(admin.bale_id)
            await send_message(bot, admin.bale_id, "???? ???? ??!", reply_markup=inline_keyboard([
                [{"text": "?? ?????? ???????", "callback_data": "admin_packages"}],
            ]))
    except Exception as e:
        await send_message(bot, admin.bale_id, f"???: {e}\n????: 1405/05/20 16:00")

async def cancel_reservation(bot, admin, res_id, db):
    res = db.query(Reservation).filter(Reservation.id == res_id).first()
    if not res:
        return
    res.status = ReservationStatus.cancelled
    if res.time_slot:
        res.time_slot.is_available = True
    db.commit()
    await send_message(bot, admin.bale_id, "???? ??? ??.")

async def prompt_reschedule(bot, admin, res_id, db):
    set_state(admin.bale_id, {"step": "admin_reschedule", "res_id": res_id})
    slots = db.query(TimeSlot).filter(TimeSlot.is_available == True, TimeSlot.date > datetime.now()).limit(8).all()
    if not slots:
        await send_message(bot, admin.bale_id, "???? ????? ????? ????.")
        return
    buttons = [[{"text": s.date.strftime("%Y/%m/%d — %H:%M"), "callback_data": f"admin_newslot_{s.id}"}] for s in slots]
    await send_message(bot, admin.bale_id, "???? ???? ?? ?????? ??:", reply_markup=inline_keyboard(buttons))

async def reschedule_reservation(bot, admin, new_slot_id, db):
    state = get_state(admin.bale_id)
    res_id = state.get("res_id")
    clear_state(admin.bale_id)
    res = db.query(Reservation).filter(Reservation.id == res_id).first()
    new_slot = db.query(TimeSlot).filter(TimeSlot.id == new_slot_id).first()
    if not res or not new_slot:
        return
    if res.time_slot:
        res.time_slot.is_available = True
    res.time_slot_id = new_slot.id
    res.status = ReservationStatus.rescheduled
    new_slot.is_available = False
    db.commit()
    new_date = new_slot.date.strftime("%Y/%m/%d — %H:%M")
    await send_message(bot, res.user.bale_id, f"???? ??? ????? ??!\n{new_date}")
    await send_message(bot, admin.bale_id, f"???? ????? ?? ?? {new_date}")

async def show_tests_management(bot, admin, db):
    tests = db.query(Test).order_by(Test.order).all()
    if not tests:
        await send_message(bot, admin.bale_id, "??? ???? ??? ????.", reply_markup=inline_keyboard([
            [{"text": "? ?????? ??? ????", "callback_data": "admin_new_test"}],
        ]))
        return
    buttons = [[{"text": f"{'?' if t.is_active else '?'} {t.title}", "callback_data": f"admin_test_menu_{t.id}"}] for t in tests]
    buttons.append([{"text": "? ?????? ??? ????", "callback_data": "admin_new_test"}])
    await send_message(bot, admin.bale_id, "?? ?????? ??????:", reply_markup=inline_keyboard(buttons))

async def show_test_menu(bot, admin, test_id, db):
    test = db.query(Test).filter(Test.id == test_id).first()
    if not test:
        return
    result_count = db.query(func.count(TestResult.id)).filter(TestResult.test_id == test_id).scalar()
    atype = getattr(test, 'analysis_type', 'numeric') or 'numeric'
    tone = getattr(test, 'tone', 'friendly') or 'friendly'
    atype_fa = {"numeric": "????", "typology": "????", "combined": "??????"}.get(atype, atype)
    tone_fa = {"friendly": "?? ?????", "formal": "?? ????"}.get(tone, tone)
    msg = (
        f"?? {test.title}\n?????????????\n"
        f"{'??????' if test.is_free else f'{test.price:,} ?????'}\n"
        f"??? ?????: {atype_fa} | ???: {tone_fa}\n"
        f"{len(test.questions)} ???? | {result_count} ??? ????? ???\n"
        f"?????: {'? ????' if test.is_active else '? ???????'}"
    )
    await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ????/???????", "callback_data": f"admin_toggle_test_{test_id}"},
         {"text": "?? ??? ???", "callback_data": f"admin_confirm_delete_test_{test_id}"}],
        [{"text": "?? ?????", "callback_data": f"admin_test_results_{test_id}"}],
        [{"text": "?? ??????", "callback_data": "admin_tests"}],
    ]))

async def confirm_delete_test(bot, admin, test_id, db):
    test = db.query(Test).filter(Test.id == test_id).first()
    if not test:
        return
    result_count = db.query(func.count(TestResult.id)).filter(TestResult.test_id == test_id).scalar()
    await send_message(bot, admin.bale_id,
        f"??? ???????\n??? {test.title} ?? {len(test.questions)} ???? ? {result_count} ????? ??? ??????.",
        reply_markup=inline_keyboard([
            [{"text": "??? ??? ??", "callback_data": f"admin_delete_test_{test_id}"},
             {"text": "??????", "callback_data": f"admin_test_menu_{test_id}"}],
        ]))

async def prompt_new_test(bot, admin):
    set_state(admin.bale_id, {"step": "admin_new_test_type"})
    await send_message(bot, admin.bale_id,
        "??? ???? - ????? 1\n??? ????? ?? ?????? ??:",
        reply_markup=inline_keyboard([
            [{"text": "?? ???? (??? ??????)", "callback_data": "admin_test_atype_numeric"}],
            [{"text": "?? ???? / ?????", "callback_data": "admin_test_atype_typology"}],
            [{"text": "?? ?????? (?????? + ???)", "callback_data": "admin_test_atype_combined"}],
        ]))

async def set_new_test_type(bot, admin, atype):
    set_state(admin.bale_id, {"step": "admin_new_test_tone", "analysis_type": atype})
    labels = {"numeric": "????", "typology": "????", "combined": "??????"}
    await send_message(bot, admin.bale_id,
        f"???: {labels.get(atype,'')}\n????? 2: ??? ??? ?? ?????? ??:",
        reply_markup=inline_keyboard([
            [{"text": "?? ?????", "callback_data": "admin_test_tone_friendly"}],
            [{"text": "?? ????", "callback_data": "admin_test_tone_formal"}],
        ]))

async def set_new_test_tone(bot, admin, tone):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_new_test", "tone": tone})
    await send_message(bot, admin.bale_id, "????? 3: ????? ??? ?? ?????:")

async def create_new_test(bot, admin, title, db):
    state = get_state(admin.bale_id)
    analysis_type = state.get("analysis_type", "numeric")
    tone = state.get("tone", "friendly")
    clear_state(admin.bale_id)
    test = Test(title=title, is_free=True, is_active=False,
                analysis_type=analysis_type, tone=tone)
    db.add(test)
    db.commit()
    db.refresh(test)
    set_state(admin.bale_id, {"step": "admin_add_question", "test_id": test.id, "q_index": 1})
    await send_message(bot, admin.bale_id,
        f"??? {title} ????? ??.\n???? 1 ?? ?????:\n(???? ???? ?? /done ???)")

async def add_question(bot, admin, text, db):
    state = get_state(admin.bale_id)
    test_id = state.get("test_id")
    if not test_id:
        await send_message(bot, admin.bale_id, "???: ??? ???? ???.")
        clear_state(admin.bale_id)
        return
    if text == "/done":
        test = db.query(Test).filter(Test.id == test_id).first()
        set_state(admin.bale_id, {"step": "admin_add_score_range", "test_id": test_id, "range_index": 1})
        await send_message(bot, admin.bale_id,
            f"{len(test.questions)} ???? ????? ??.\n\n???? ???????? ?????? ?? ???? ??:\n"
            "????: ????? | ?????? | ????? | ??? ?????\n\n"
            "????:\n0 | 10 | ??? ????? | ???? ?? ????? ?????\n"
            "10 | 20 | ??? ????? | ?????? ???? ???? ???\n\n"
            "???? ???? ?? /done ???")
        return
    q = Question(test_id=test_id, text=text, order=state.get("q_index", 1)-1)
    db.add(q)
    db.commit()
    db.refresh(q)
    set_state(admin.bale_id, {**state, "step": "admin_add_question_analysis", "question_id": q.id})
    await send_message(bot, admin.bale_id,
        "???? ????? ??.\n????? ??? ???? ?? ?????:\n(???? ?? ???? /skip ???)")

async def save_question_analysis(bot, admin, text, db):
    state = get_state(admin.bale_id)
    question_id = state.get("question_id")
    if question_id and text != "/skip":
        q = db.query(Question).filter(Question.id == question_id).first()
        if q:
            q.analysis = text
            db.commit()
    set_state(admin.bale_id, {**state, "step": "admin_add_options", "opt_index": 1})
    await send_message(bot, admin.bale_id,
        "???????? ?? ???? ??:\n"
        "????: ??? | ??????\n"
        "????: ??? | ???\n"
        "??????: ??? | ?????? | ???\n\n"
        "???? ???? ?? /next ???")

async def add_option(bot, admin, text, db):
    state = get_state(admin.bale_id)
    if text == "/next":
        q_index = state["q_index"] + 1
        set_state(admin.bale_id, {**state, "step": "admin_add_question", "q_index": q_index})
        await send_message(bot, admin.bale_id,
            f"???? {q_index} ?? ?????:\n(???? ???? ?? /done ???)")
        return
    try:
        parts = [p.strip() for p in text.split("|")]
        opt_text = parts[0]
        score = 0.0
        type_label = None
        if len(parts) >= 2:
            try:
                score = float(parts[1])
            except:
                type_label = parts[1]
        if len(parts) >= 3:
            type_label = parts[2]
        opt = Option(question_id=state["question_id"], text=opt_text,
                     score=score, type_label=type_label, order=state["opt_index"]-1)
        db.add(opt)
        db.commit()
        set_state(admin.bale_id, {**state, "opt_index": state["opt_index"] + 1})
        label_info = f" | ???: {type_label}" if type_label else ""
        await send_message(bot, admin.bale_id,
            f"????? ????? ??: {opt_text} (??????: {score}{label_info})\n????? ???? ?? /next")
    except:
        await send_message(bot, admin.bale_id, "???? ??????.")

async def add_score_range(bot, admin, text, db):
    state = get_state(admin.bale_id)
    test_id = state.get("test_id")
    if text == "/done":
        clear_state(admin.bale_id)
        test = db.query(Test).filter(Test.id == test_id).first()
        count = db.query(ScoreRange).filter(ScoreRange.test_id == test_id).count()
        await send_message(bot, admin.bale_id,
            f"??? {test.title} ?? {len(test.questions)} ???? ? {count} ???? ????? ??.",
            reply_markup=inline_keyboard([
                [{"text": "?? ???? ???? ???", "callback_data": f"admin_toggle_test_{test_id}"}],
                [{"text": "?? ?????? ??????", "callback_data": "admin_tests"}],
            ]))
        return
    try:
        parts = [p.strip() for p in text.split("|")]
        if len(parts) < 4:
            raise ValueError("???? ????")
        min_score = float(parts[0])
        max_score = float(parts[1])
        title = parts[2]
        analysis = "|".join(parts[3:]).strip()
        sr = ScoreRange(test_id=test_id, min_score=min_score, max_score=max_score,
                        title=title, analysis=analysis)
        db.add(sr)
        db.commit()
        idx = state.get("range_index", 1)
        set_state(admin.bale_id, {**state, "range_index": idx + 1})
        await send_message(bot, admin.bale_id,
            f"???? {idx}: {min_score} ?? {max_score} - {title}\n???? ???? ?? /done")
    except Exception as e:
        await send_message(bot, admin.bale_id,
            f"???? ??????: {e}\n????: 0 | 10 | ??? ????? | ??? ?????")

async def toggle_test(bot, admin, test_id, db):
    test = db.query(Test).filter(Test.id == test_id).first()
    if not test:
        return
    test.is_active = not test.is_active
    db.commit()
    await send_message(bot, admin.bale_id, f"??? {'????' if test.is_active else '???????'} ??.")

async def delete_test(bot, admin, test_id, db):
    test = db.query(Test).filter(Test.id == test_id).first()
    if not test:
        return
    title = test.title
    db.delete(test)
    db.commit()
    await send_message(bot, admin.bale_id, f"??? {title} ??? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ?????? ??????", "callback_data": "admin_tests"}],
    ]))

async def show_products_management(bot, admin, db):
    products = db.query(Product).order_by(Product.order).all()
    buttons = [[{"text": f"{'?' if p.is_active else '?'} {p.title}", "callback_data": f"admin_product_menu_{p.id}"}] for p in products]
    buttons.append([{"text": "? ?????? ?????", "callback_data": "admin_new_product"}])
    await send_message(bot, admin.bale_id, "?? ?????? ???????:", reply_markup=inline_keyboard(buttons))

async def show_product_menu(bot, admin, product_id, db):
    p = db.query(Product).filter(Product.id == product_id).first()
    if not p:
        return
    msg = (
        f"{p.title}\n?????????????\n"
        f"{'??????' if p.price == 0 else f'{p.price:,} ?????'}\n"
        f"{'?? ' + p.link if p.link else ''}\n"
        f"?????: {'? ????' if p.is_active else '? ???????'}"
    )
    await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ????/???????", "callback_data": f"admin_toggle_product_{product_id}"},
         {"text": "?? ??? ?????", "callback_data": f"admin_delete_product_{product_id}"}],
        [{"text": "?? ??????", "callback_data": "admin_products"}],
    ]))

async def toggle_product(bot, admin, product_id, db):
    p = db.query(Product).filter(Product.id == product_id).first()
    if not p:
        return
    p.is_active = not p.is_active
    db.commit()
    await send_message(bot, admin.bale_id, f"????? {'????' if p.is_active else '???????'} ??.")

async def delete_product(bot, admin, product_id, db):
    p = db.query(Product).filter(Product.id == product_id).first()
    if not p:
        return
    title = p.title
    db.delete(p)
    db.commit()
    await send_message(bot, admin.bale_id, f"????? {title} ??? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ?????? ???????", "callback_data": "admin_products"}],
    ]))

async def prompt_new_product(bot, admin):
    set_state(admin.bale_id, {"step": "admin_new_product_type"})
    await send_message(bot, admin.bale_id, "??? ?????:", reply_markup=inline_keyboard([
        [{"text": "?? ????", "callback_data": "admin_ptype_course"},
         {"text": "?? ??????", "callback_data": "admin_ptype_podcast"}],
        [{"text": "?? ??????", "callback_data": "admin_ptype_webinar"},
         {"text": "?? ?????", "callback_data": "admin_ptype_channel"}],
    ]))

async def set_product_type(bot, admin, ptype):
    set_state(admin.bale_id, {"step": "admin_new_product_title", "ptype": ptype})
    await send_message(bot, admin.bale_id, "????? ?????:")

async def set_product_title(bot, admin, title):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_new_product_price", "title": title})
    await send_message(bot, admin.bale_id, "???? (?????):\n(???? ??????: 0)")

async def set_product_price(bot, admin, price_text, db):
    state = get_state(admin.bale_id)
    try:
        price = int(price_text.replace(",", ""))
        set_state(admin.bale_id, {**state, "step": "admin_new_product_capacity", "price": price})
        await send_message(bot, admin.bale_id, "?????:\n(???? ???????: 0)")
    except:
        await send_message(bot, admin.bale_id, "???? ?????? ???.")

async def set_product_capacity(bot, admin, cap_text, db):
    state = get_state(admin.bale_id)
    try:
        capacity = int(cap_text.strip())
        product = Product(title=state["title"], type=ProductType(state["ptype"]),
                          price=state["price"], capacity=capacity, is_active=True)
        db.add(product)
        db.commit()
        db.refresh(product)
        if state.get("price", 0) > 0:
            await prompt_installment_question(bot, admin, "product", product.id, state.get("price", 0))
        else:
            set_state(admin.bale_id, {"step": "admin_product_link", "product_id": product.id})
            await send_message(bot, admin.bale_id, "???? ?? ?????:\n(???? ??: /skip)")
    except:
        await send_message(bot, admin.bale_id, "??? ?????? ???.")

async def set_product_link(bot, admin, text, db):
    state = get_state(admin.bale_id)
    clear_state(admin.bale_id)
    if text != "/skip":
        product = db.query(Product).filter(Product.id == state["product_id"]).first()
        if product:
            product.link = text
            db.commit()
    await send_message(bot, admin.bale_id, "????? ?? ????? ??????? ???.", reply_markup=inline_keyboard([
        [{"text": "?? ?????? ???????", "callback_data": "admin_products"}],
    ]))

async def show_pending_reviews(bot, admin, db):
    from models.review import Review
    reviews = db.query(Review).filter(Review.is_approved == False).limit(5).all()
    if not reviews:
        await send_message(bot, admin.bale_id, "??? ?? ??????? ???? ?????.")
        return
    for r in reviews:
        subject = r.test.title if r.test else (r.product.title if r.product else "??????")
        msg = f"? {r.score}/5\n?? {r.user.full_name}\n?? {subject}\n?? {r.text or '???? ???'}"
        await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
            [{"text": "?????", "callback_data": f"admin_approve_review_{r.id}"},
             {"text": "??", "callback_data": f"admin_reject_review_{r.id}"}],
        ]))

async def approve_review(bot, admin, review_id, db):
    from models.review import Review
    r = db.query(Review).filter(Review.id == review_id).first()
    if r:
        r.is_approved = True
        r.allow_publish = True
        db.commit()
    await send_message(bot, admin.bale_id, "??? ????? ??.")

async def reject_review(bot, admin, review_id, db):
    from models.review import Review
    r = db.query(Review).filter(Review.id == review_id).first()
    if r:
        db.delete(r)
        db.commit()
    await send_message(bot, admin.bale_id, "??? ??? ??.")

async def show_financial_report(bot, admin, db):
    today = datetime.now()
    week_ago = today - timedelta(days=7)
    month_ago = today - timedelta(days=30)
    def income(since=None):
        q = db.query(func.sum(Payment.amount)).filter(Payment.status == PaymentStatus.approved)
        if since:
            q = q.filter(Payment.updated_at >= since)
        return q.scalar() or 0
    msg = (
        "????? ????\n?????????????\n"
        f"?????: {income(today.replace(hour=0,minute=0)):,} ?????\n"
        f"??? ????: {income(week_ago):,} ?????\n"
        f"??? ???: {income(month_ago):,} ?????\n"
        f"?????????????\n??: {income():,} ?????"
    )
    await send_message(bot, admin.bale_id, msg)

async def show_settings(bot, admin, db):
    card_info = ""
    if bot_settings.get("card_number"):
        card_info = f"\n????? ????: {bot_settings['card_number']}"
        if bot_settings.get("card_owner"):
            card_info += f"\n?? ???: {bot_settings['card_owner']}"
    await send_message(bot, admin.bale_id, f"??????? ????:{card_info}", reply_markup=inline_keyboard([
        [{"text": "?? ?????? ???? ???????????", "callback_data": "admin_edit_welcome"}],
        [{"text": "?? ?????? ????? ????", "callback_data": "admin_edit_card"}],
        [{"text": "?? ?????? ??? ???? ????", "callback_data": "admin_edit_card_owner"}],
    ]))

async def prompt_edit_welcome(bot, admin):
    set_state(admin.bale_id, {"step": "admin_edit_welcome"})
    await send_message(bot, admin.bale_id,
        f"???? ????:\n\n{bot_settings.get('welcome_message','')}\n\n???? ???? ?? ?????:")

async def save_welcome_message(bot, admin, text):
    clear_state(admin.bale_id)
    bot_settings["welcome_message"] = text
    await send_message(bot, admin.bale_id, "???? ????? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ????? ?????", "callback_data": "admin_discounts"},
         {"text": "?? ???????", "callback_data": "admin_settings"}],
    ]))

async def prompt_edit_card(bot, admin):
    set_state(admin.bale_id, {"step": "admin_edit_card"})
    await send_message(bot, admin.bale_id,
        f"???? ????: {bot_settings.get('card_number','??? ????')}\n????? ???? ????:")

async def save_card_number(bot, admin, text):
    clear_state(admin.bale_id)
    bot_settings["card_number"] = text.strip()
    await send_message(bot, admin.bale_id, f"???? ????? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ????? ?????", "callback_data": "admin_discounts"},
         {"text": "?? ???????", "callback_data": "admin_settings"}],
    ]))

async def prompt_edit_card_owner(bot, admin):
    set_state(admin.bale_id, {"step": "admin_edit_card_owner"})
    await send_message(bot, admin.bale_id,
        f"??? ????: {bot_settings.get('card_owner','??? ????')}\n??? ????:")

async def save_card_owner(bot, admin, text):
    clear_state(admin.bale_id)
    bot_settings["card_owner"] = text.strip()
    await send_message(bot, admin.bale_id, f"??? ????? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ????? ?????", "callback_data": "admin_discounts"},
         {"text": "?? ???????", "callback_data": "admin_settings"}],
    ]))

async def show_packages_management(bot, admin, db):
    packages = db.query(CoachingPackage).order_by(CoachingPackage.order).all()
    buttons = [[{"text": f"{'?' if p.is_active else '?'} {p.title}", "callback_data": f"admin_pkg_menu_{p.id}"}] for p in packages]
    buttons.append([{"text": "? ?????? ???? ????", "callback_data": "admin_new_package"}])
    await send_message(bot, admin.bale_id, "?? ?????? ???????? ??????:", reply_markup=inline_keyboard(buttons))

async def show_package_menu(bot, admin, pkg_id, db):
    p = db.query(CoachingPackage).filter(CoachingPackage.id == pkg_id).first()
    if not p:
        return
    msg = (
        f"{p.title}\n?????????????\n"
        f"{p.session_count} ???? - {p.session_duration} ?????\n"
        f"{p.price:,} ?????\n"
        f"?????: {'? ????' if p.is_active else '? ???????'}"
    )
    await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ????/???????", "callback_data": f"admin_toggle_pkg_{pkg_id}"},
         {"text": "?? ???", "callback_data": f"admin_delete_pkg_{pkg_id}"}],
        [{"text": "?? ??????", "callback_data": "admin_packages"}],
    ]))

async def toggle_package(bot, admin, pkg_id, db):
    p = db.query(CoachingPackage).filter(CoachingPackage.id == pkg_id).first()
    if p:
        p.is_active = not p.is_active
        db.commit()
        await send_message(bot, admin.bale_id, f"???? {'????' if p.is_active else '???????'} ??.")

async def delete_package(bot, admin, pkg_id, db):
    p = db.query(CoachingPackage).filter(CoachingPackage.id == pkg_id).first()
    if p:
        title = p.title
        db.delete(p)
        db.commit()
        await send_message(bot, admin.bale_id, f"???? {title} ??? ??.")

async def prompt_new_package(bot, admin):
    set_state(admin.bale_id, {"step": "admin_new_package_title"})
    await send_message(bot, admin.bale_id, "????? ???? ????:")

async def set_package_title(bot, admin, text):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_new_package_sessions", "pkg_title": text})
    await send_message(bot, admin.bale_id, "????? ?????:")

async def set_package_sessions(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        sessions = int(text)
        set_state(admin.bale_id, {**state, "step": "admin_new_package_duration", "pkg_sessions": sessions})
        await send_message(bot, admin.bale_id, "??? ?? ???? (?????):")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def set_package_duration(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        duration = int(text)
        set_state(admin.bale_id, {**state, "step": "admin_new_package_price", "pkg_duration": duration})
        await send_message(bot, admin.bale_id, "???? ?? (?????):")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def set_package_price(bot, admin, text, db):
    state = get_state(admin.bale_id)
    try:
        price = int(text.replace(",", ""))
        count = db.query(CoachingPackage).count()
        pkg = CoachingPackage(title=state["pkg_title"], session_count=state["pkg_sessions"],
                              session_duration=state["pkg_duration"], price=price,
                              is_active=True, order=count)
        db.add(pkg)
        db.commit()
        db.refresh(pkg)
        set_state(admin.bale_id, {**state, "pkg_id_new": pkg.id})
        if price > 0:
            await prompt_installment_question(bot, admin, "package", pkg.id, price)
        else:
            clear_state(admin.bale_id)
            await send_message(bot, admin.bale_id, "???? ????? ??!", reply_markup=inline_keyboard([
                [{"text": "?????? ???????", "callback_data": "admin_packages"}],
            ]))
    except:
        await send_message(bot, admin.bale_id, "???? ?????? ???.")

CONTENT_TYPE_LABELS = {"review": "? ????? ???????", "intro": "?? ????? ??", "faq": "? ?????? ??????"}
MEDIA_TYPE_LABELS = {
    "text": "?? ???", "photo": "?? ???", "video": "?? ?????",
    "voice": "???", "audio": "?? ???? ????", "document": "????", "link": "?? ????"
}

async def show_static_content_menu(bot, admin, db):
    await send_message(bot, admin.bale_id, "?? ?????? ?????? ?????:", reply_markup=inline_keyboard([
        [{"text": "? ????? ???????", "callback_data": "admin_sc_list_review"},
         {"text": "?? ????? ??", "callback_data": "admin_sc_list_intro"}],
        [{"text": "? ?????? ??????", "callback_data": "admin_sc_list_faq"}],
        [{"text": "?? ??????", "callback_data": "admin_dashboard"}],
    ]))

async def show_sc_list(bot, admin, ctype, db):
    items = db.query(StaticContent).filter(StaticContent.content_type == ctype).order_by(StaticContent.order).all()
    label = CONTENT_TYPE_LABELS.get(ctype, ctype)
    buttons = []
    for item in items:
        preview = item.question or item.text or item.link_url or f"???? {item.id}"
        preview = (preview[:28] + "...") if len(preview) > 28 else preview
        status = "?" if item.is_active else "?"
        mtype = MEDIA_TYPE_LABELS.get(item.media_type, "")
        buttons.append([{"text": f"{status} {mtype} - {preview}", "callback_data": f"admin_sc_item_{item.id}"}])
    buttons.append([{"text": "? ??????", "callback_data": f"admin_sc_add_{ctype}"}])
    buttons.append([{"text": "?? ??????", "callback_data": "admin_static_content"}])
    await send_message(bot, admin.bale_id, f"{label} - {len(items)} ????:", reply_markup=inline_keyboard(buttons))

async def show_sc_item(bot, admin, item_id, db):
    item = db.query(StaticContent).filter(StaticContent.id == item_id).first()
    if not item:
        return
    label = CONTENT_TYPE_LABELS.get(item.content_type, "")
    mtype = MEDIA_TYPE_LABELS.get(item.media_type, "")
    preview = item.question or item.text or item.link_url or "-"
    preview = (preview[:50] + "...") if len(preview) > 50 else preview
    msg = f"{label} - {mtype}\n{preview}\n?????: {'? ????' if item.is_active else '? ???????'}"
    await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ????/???????", "callback_data": f"admin_sc_toggle_{item_id}"},
         {"text": "?? ???", "callback_data": f"admin_sc_delete_{item_id}"}],
        [{"text": "?? ??????", "callback_data": f"admin_sc_list_{item.content_type}"}],
    ]))

async def toggle_sc_item(bot, admin, item_id, db):
    item = db.query(StaticContent).filter(StaticContent.id == item_id).first()
    if item:
        item.is_active = not item.is_active
        db.commit()
        await send_message(bot, admin.bale_id, f"{'????' if item.is_active else '???????'} ??.")

async def delete_sc_item(bot, admin, item_id, db):
    item = db.query(StaticContent).filter(StaticContent.id == item_id).first()
    if item:
        ctype = item.content_type
        db.delete(item)
        db.commit()
        await send_message(bot, admin.bale_id, "??? ??.", reply_markup=inline_keyboard([
            [{"text": "?? ??????", "callback_data": f"admin_sc_list_{ctype}"}],
        ]))

async def prompt_add_sc(bot, admin, ctype):
    label = CONTENT_TYPE_LABELS.get(ctype, ctype)
    set_state(admin.bale_id, {"step": "admin_sc_choose_media", "sc_ctype": ctype})
    await send_message(bot, admin.bale_id,
        f"?????? ?? {label}\n??? ????? ?? ?????? ??:",
        reply_markup=inline_keyboard([
            [{"text": "?? ???", "callback_data": "admin_sc_media_text"},
             {"text": "?? ????", "callback_data": "admin_sc_media_link"}],
            [{"text": "?? ???", "callback_data": "admin_sc_media_photo"},
             {"text": "?? ?????", "callback_data": "admin_sc_media_video"}],
            [{"text": "?? ??? (voice)", "callback_data": "admin_sc_media_voice"},
             {"text": "?? ???? ????", "callback_data": "admin_sc_media_audio"}],
            [{"text": "?? ???? / ???", "callback_data": "admin_sc_media_document"}],
        ]))

async def sc_set_media_type(bot, admin, mtype):
    state = get_state(admin.bale_id)
    ctype = state.get("sc_ctype", "review")
    if ctype == "faq":
        set_state(admin.bale_id, {**state, "step": "admin_sc_faq_question", "sc_mtype": mtype})
        await send_message(bot, admin.bale_id, "??? ???? ?? ?????:")
    elif mtype == "text":
        set_state(admin.bale_id, {**state, "step": "admin_sc_text_input", "sc_mtype": mtype})
        await send_message(bot, admin.bale_id, "??? ?? ?????:")
    elif mtype == "link":
        set_state(admin.bale_id, {**state, "step": "admin_sc_link_url", "sc_mtype": mtype})
        await send_message(bot, admin.bale_id, "???? ?? ???? ??:")
    else:
        set_state(admin.bale_id, {**state, "step": "admin_sc_await_file", "sc_mtype": mtype})
        hint = {"photo": "?? ???", "video": "?? ?????", "voice": "???? ????",
                "audio": "?? ???? ????", "document": "????"}.get(mtype, "????")
        await send_message(bot, admin.bale_id, f"{hint} ?? ?????:")

async def sc_save_text(bot, admin, text, db):
    state = get_state(admin.bale_id)
    clear_state(admin.bale_id)
    ctype = state.get("sc_ctype")
    count = db.query(StaticContent).filter(StaticContent.content_type == ctype).count()
    db.add(StaticContent(content_type=ctype, media_type="text", text=text, is_active=True, order=count))
    db.commit()
    await send_message(bot, admin.bale_id, "??? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ??????", "callback_data": f"admin_sc_list_{ctype}"}],
    ]))

async def sc_save_faq_question(bot, admin, text):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_sc_faq_answer", "sc_question": text})
    await send_message(bot, admin.bale_id, "???? ??? ???? ?? ?????:")

async def sc_save_faq_answer(bot, admin, text, db):
    state = get_state(admin.bale_id)
    clear_state(admin.bale_id)
    count = db.query(StaticContent).filter(StaticContent.content_type == "faq").count()
    db.add(StaticContent(content_type="faq", media_type="text",
                         question=state.get("sc_question"), text=text, is_active=True, order=count))
    db.commit()
    await send_message(bot, admin.bale_id, "???? ? ???? ??? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ??????", "callback_data": "admin_sc_list_faq"}],
    ]))

async def sc_save_link_url(bot, admin, url):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_sc_link_label", "sc_link_url": url})
    await send_message(bot, admin.bale_id, "???? ???? ?? ?????:")

async def sc_save_link(bot, admin, label, db):
    state = get_state(admin.bale_id)
    clear_state(admin.bale_id)
    ctype = state.get("sc_ctype")
    count = db.query(StaticContent).filter(StaticContent.content_type == ctype).count()
    db.add(StaticContent(content_type=ctype, media_type="link",
                         link_url=state.get("sc_link_url"), link_label=label, is_active=True, order=count))
    db.commit()
    await send_message(bot, admin.bale_id, "???? ??? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ??????", "callback_data": f"admin_sc_list_{ctype}"}],
    ]))

async def sc_save_file(bot, admin, file_id, caption, db):
    state = get_state(admin.bale_id)
    clear_state(admin.bale_id)
    ctype = state.get("sc_ctype")
    mtype = state.get("sc_mtype", "document")
    count = db.query(StaticContent).filter(StaticContent.content_type == ctype).count()
    db.add(StaticContent(content_type=ctype, media_type=mtype, file_id=file_id,
                         text=caption or None, is_active=True, order=count))
    db.commit()
    label = CONTENT_TYPE_LABELS.get(ctype, "")
    await send_message(bot, admin.bale_id, f"???? ?? {label} ????? ??.", reply_markup=inline_keyboard([
        [{"text": "?? ??????", "callback_data": f"admin_sc_list_{ctype}"}],
    ]))

async def show_discounts(bot, admin, db):
    codes = db.query(DiscountCode).order_by(DiscountCode.created_at.desc()).all()
    buttons = []
    for c in codes:
        status = "?" if c.is_active else "?"
        uses = f"{c.used_count}" + (f"/{c.max_uses}" if c.max_uses > 0 else "")
        buttons.append([{"text": f"{status} {c.code} — {c.amount:,} ????? | {uses} ???", "callback_data": f"admin_discount_{c.id}"}])
    buttons.append([{"text": "? ?? ????", "callback_data": "admin_new_discount"}])
    await send_message(bot, admin.bale_id, f"?? ????? ????? ({len(codes)} ??):", reply_markup=inline_keyboard(buttons))

async def prompt_new_discount(bot, admin):
    set_state(admin.bale_id, {"step": "admin_discount_code"})
    await send_message(bot, admin.bale_id, "?? ????? ?? ?????:\n(????: SPRING10)")

async def discount_set_code(bot, admin, text):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_discount_amount", "disc_code": text.strip().upper()})
    await send_message(bot, admin.bale_id, "???? ????? (?????):\n(????: 100000)")

async def discount_set_amount(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        amount = int(text.replace(",", ""))
        set_state(admin.bale_id, {**state, "step": "admin_discount_maxuses", "disc_amount": amount})
        await send_message(bot, admin.bale_id, "?????? ????? ???????:\n(0 ???? ???????)")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def discount_set_maxuses(bot, admin, text, db):
    state = get_state(admin.bale_id)
    clear_state(admin.bale_id)
    try:
        max_uses = int(text.strip())
        dc = DiscountCode(code=state["disc_code"], amount=state["disc_amount"], max_uses=max_uses)
        db.add(dc)
        db.commit()
        await send_message(bot, admin.bale_id,
            f"? ?? {state['disc_code']} ?? ????? {state['disc_amount']:,} ????? ????? ??.",
            reply_markup=inline_keyboard([[{"text": "?? ????? ?????", "callback_data": "admin_discounts"}]]))
    except Exception as e:
        await send_message(bot, admin.bale_id, f"???: {e}")

async def show_discount_item(bot, admin, dc_id, db):
    dc = db.query(DiscountCode).filter(DiscountCode.id == dc_id).first()
    if not dc:
        return
    uses = f"{dc.used_count}" + (f" ?? {dc.max_uses}" if dc.max_uses > 0 else " (???????)")
    msg = (f"?? {dc.code}\n?????????????\n"
           f"?? ?????: {dc.amount:,} ?????\n"
           f"?? ???????: {uses}\n"
           f"?????: {'? ????' if dc.is_active else '? ???????'}")
    await send_message(bot, admin.bale_id, msg, reply_markup=inline_keyboard([
        [{"text": "?? ????/???????", "callback_data": f"admin_toggle_discount_{dc_id}"},
         {"text": "?? ???", "callback_data": f"admin_delete_discount_{dc_id}"}],
        [{"text": "?? ??????", "callback_data": "admin_discounts"}],
    ]))

async def toggle_discount(bot, admin, dc_id, db):
    dc = db.query(DiscountCode).filter(DiscountCode.id == dc_id).first()
    if dc:
        dc.is_active = not dc.is_active
        db.commit()
        await send_message(bot, admin.bale_id, f"{'? ????' if dc.is_active else '? ???????'} ??.")

async def delete_discount(bot, admin, dc_id, db):
    dc = db.query(DiscountCode).filter(DiscountCode.id == dc_id).first()
    if dc:
        db.delete(dc)
        db.commit()
        await send_message(bot, admin.bale_id, "?? ?? ??? ??.", reply_markup=inline_keyboard([
            [{"text": "?? ????? ?????", "callback_data": "admin_discounts"}],
        ]))

async def prompt_reject_payment(bot, admin, payment_id):
    from bot.handlers.payment import prompt_reject_reason
    await prompt_reject_reason(bot, admin.bale_id, payment_id)

async def prompt_installment_setup(bot, admin, entity_type, entity_id, price):
    """???? ????? ????? ???? ?????/????"""
    set_state(admin.bale_id, {
        "step": "admin_installment_first",
        "inst_type": entity_type,
        "inst_id": entity_id,
        "inst_price": price,
    })
    await send_message(bot, admin.bale_id,
        f"????? ???? ??? ????\n???? ??: {price:,} ?????\n\n???? ?????? ????? (?????):")

async def installment_set_first(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        first = int(text.replace(",", ""))
        set_state(admin.bale_id, {**state, "step": "admin_installment_count", "inst_first": first})
        await send_message(bot, admin.bale_id, "????? ????? ??? ?? ?????? ?????:")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def installment_set_count(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        count = int(text.strip())
        set_state(admin.bale_id, {**state, "step": "admin_installment_details",
                                   "inst_count": count, "inst_current": 1, "inst_data": []})
        await send_message(bot, admin.bale_id,
            f"??? 1 ?? {count}\n????: ???? | ?????\n????: 500000 | 1405/06/01")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def installment_add_detail(bot, admin, text, db):
    state = get_state(admin.bale_id)
    count = state.get("inst_count", 1)
    current = state.get("inst_current", 1)
    inst_data = state.get("inst_data", [])
    try:
        import jdatetime
        parts = [p.strip() for p in text.split("|")]
        amount = int(parts[0].replace(",", ""))
        jalali_dt = jdatetime.datetime.strptime(parts[1].strip(), "%Y/%m/%d")
        due_date = jalali_dt.togregorian()
        inst_data.append({"amount": amount, "due_date": due_date, "number": current})
        if current < count:
            set_state(admin.bale_id, {**state, "inst_current": current+1, "inst_data": inst_data})
            await send_message(bot, admin.bale_id,
                f"??? {current} ??? ??.\n??? {current+1} ?? {count}\n????: ???? | ?????")
        else:
            # ????? ???
            clear_state(admin.bale_id)
            plan = InstallmentPlan(
                product_id=state.get("inst_id") if state.get("inst_type") == "product" else None,
                package_id=state.get("inst_id") if state.get("inst_type") == "package" else None,
                first_payment=state.get("inst_first"),
                total_count=count,
            )
            db.add(plan)
            db.commit()
            db.refresh(plan)
            for d in inst_data:
                inst = Installment(
                    plan_id=plan.id, number=d["number"],
                    amount=d["amount"], due_date=d["due_date"]
                )
                db.add(inst)
            db.commit()
            await send_message(bot, admin.bale_id,
                f"??? ?????? ?? {count} ??? ????? ??.",
                reply_markup=inline_keyboard([[{"text": "?? ???????", "callback_data": "admin_dashboard"}]]))
    except Exception as e:
        await send_message(bot, admin.bale_id, f"???: {e}\n????: 500000 | 1405/06/01")

async def approve_installment_admin(bot, admin, ui_id, db):
    from bot.handlers.installment import approve_installment
    await approve_installment(bot, ui_id, db)
    await send_message(bot, admin.bale_id, "??? ????? ??.")

async def reject_installment_admin(bot, admin, ui_id):
    from bot.handlers.installment import prompt_reject_installment
    await prompt_reject_installment(bot, admin.bale_id, ui_id)


# ==================== ????? ====================

async def prompt_installment_question(bot, admin, entity_type, entity_id, price):
    set_state(admin.bale_id, {
        "step": "admin_installment_yn",
        "inst_type": entity_type,
        "inst_id": entity_id,
        "inst_price": price,
    })
    await send_message(bot, admin.bale_id,
        f"??? ??? ???? ????? ?????",
        reply_markup=inline_keyboard([
            [{"text": "???? ????? ????", "callback_data": "admin_inst_yes"}],
            [{"text": "???? ???? ???", "callback_data": "admin_inst_no"}],
        ]))

async def handle_inst_yes(bot, admin):
    state = get_state(admin.bale_id)
    set_state(admin.bale_id, {**state, "step": "admin_inst_first_payment"})
    await send_message(bot, admin.bale_id,
        f"???? ??: {state.get('inst_price',0):,} ?????\n???? ?????? ????? (?????):")

async def handle_inst_no(bot, admin, db):
    state = get_state(admin.bale_id)
    entity_type = state.get("inst_type")
    entity_id = state.get("inst_id")
    if entity_type == "package":
        pkg = db.query(CoachingPackage).filter(CoachingPackage.id == entity_id).first()
        if pkg:
            capacity = state.get("capacity", 1)
            sessions = pkg.session_count
            total_slots = sessions * capacity
            set_state(admin.bale_id, {
                **state, "step": "admin_slot_datetime", "pkg_id": pkg.id,
                "current_session": 1, "current_person": 1,
                "collected_slots": [], "capacity": capacity,
                "total_slots": total_slots, "pkg_sessions": sessions,
            })
            await send_message(bot, admin.bale_id,
                f"???? ?????.\n??? 1 ?? {capacity}\n???? 1 ?? {sessions}\n????: 1405/05/20 16:00")
    elif entity_type == "product":
        set_state(admin.bale_id, {"step": "admin_product_link", "product_id": entity_id})
        await send_message(bot, admin.bale_id, "???? ?? ?????:\n(???? ??: /skip)")

async def inst_set_first_payment(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        first = int(text.replace(",", ""))
        set_state(admin.bale_id, {**state, "step": "admin_inst_count", "inst_first": first})
        await send_message(bot, admin.bale_id, "????? ????? (??? ?? ?????? ?????):")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def inst_set_count(bot, admin, text):
    state = get_state(admin.bale_id)
    try:
        count = int(text.strip())
        set_state(admin.bale_id, {**state, "step": "admin_inst_details",
                                   "inst_count": count, "inst_current": 1, "inst_data": []})
        await send_message(bot, admin.bale_id,
            f"??? 1 ?? {count}\n????: ???? | ?????\n????: 500000 | 1405/06/01")
    except:
        await send_message(bot, admin.bale_id, "??? ???? ??.")

async def inst_add_detail(bot, admin, text, db):
    state = get_state(admin.bale_id)
    count = state.get("inst_count", 1)
    current = state.get("inst_current", 1)
    inst_data = state.get("inst_data", [])
    try:
        import jdatetime
        parts = [p.strip() for p in text.split("|")]
        amount = int(parts[0].replace(",", ""))
        jalali_dt = jdatetime.datetime.strptime(parts[1].strip(), "%Y/%m/%d")
        due_date = jalali_dt.togregorian()
        inst_data.append({"amount": amount, "due_date": due_date, "number": current})

        if current < count:
            set_state(admin.bale_id, {**state, "inst_current": current+1, "inst_data": inst_data})
            await send_message(bot, admin.bale_id,
                f"??? {current} ??? ??.\n??? {current+1} ?? {count}\n????: ???? | ?????")
        else:
            # ????? ??? ?????
            entity_type = state.get("inst_type")
            entity_id = state.get("inst_id")
            from models.installment import InstallmentPlan, Installment
            plan = InstallmentPlan(
                product_id=entity_id if entity_type == "product" else None,
                package_id=entity_id if entity_type == "package" else None,
                first_payment=state.get("inst_first"),
                total_count=count,
            )
            db.add(plan)
            db.commit()
            db.refresh(plan)
            for d in inst_data:
                inst = Installment(plan_id=plan.id, number=d["number"],
                                   amount=d["amount"], due_date=d["due_date"])
                db.add(inst)
            db.commit()
            # ??? ???? ???? ??? ????? ?????????
            if state.get("inst_type") == "package":
                pkg_id = state.get("inst_id")
                pkg2 = db.query(CoachingPackage).filter(CoachingPackage.id == pkg_id).first()
                if pkg2:
                    capacity = state.get("capacity", 1)
                    sessions = pkg2.session_count
                    total_slots = sessions * capacity
                    set_state(admin.bale_id, {
                        **state, "step": "admin_slot_datetime", "pkg_id": pkg_id,
                        "current_session": 1, "current_person": 1,
                        "collected_slots": [], "capacity": capacity,
                        "total_slots": total_slots, "pkg_sessions": sessions,
                    })
                    await send_message(bot, admin.bale_id,
                        f"??? ?????? ????? ??.\n??? 1 ?? {capacity}\n???? 1 ?? {sessions}\n????: 1405/05/20 16:00")
                    return
            clear_state(admin.bale_id)
            await send_message(bot, admin.bale_id,
                f"??? ?????? ?? {count} ??? ????? ??.",
                reply_markup=inline_keyboard([
                    [{"text": "???????", "callback_data": "admin_dashboard"}],
                ]))
    except Exception as e:
        await send_message(bot, admin.bale_id, f"???: {e}\n????: 500000 | 1405/06/01")

async def handle_admin_action(bot, admin, data, db):
    if data == "admin_pending_payments":
        await show_pending_payments(bot, admin, db)
    elif data == "admin_reservations":
        await show_reservations(bot, admin, db)
    elif data == "admin_search_user":
        await prompt_search_user(bot, admin)
    elif data == "admin_broadcast":
        await prompt_broadcast(bot, admin)
    elif data == "admin_add_slot":
        await prompt_add_slot(bot, admin)
    elif data == "admin_dashboard":
        await show_dashboard(bot, admin, db)
    elif data.startswith("admin_slottype_"):
        await prompt_add_slot_datetime(bot, admin, data.replace("admin_slottype_", ""))
    elif data == "admin_tests":
        await show_tests_management(bot, admin, db)
    elif data == "admin_products":
        await show_products_management(bot, admin, db)
    elif data == "admin_packages":
        await show_packages_management(bot, admin, db)
    elif data == "admin_new_package":
        await prompt_new_package(bot, admin)
    elif data.startswith("admin_pkg_menu_"):
        await show_package_menu(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_toggle_pkg_"):
        await toggle_package(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_delete_pkg_"):
        await delete_package(bot, admin, int(data.split("_")[-1]), db)
    elif data == "admin_reviews":
        await show_pending_reviews(bot, admin, db)
    elif data == "admin_report":
        await show_financial_report(bot, admin, db)
    elif data == "admin_settings":
        await show_settings(bot, admin, db)
    elif data == "admin_new_test":
        await prompt_new_test(bot, admin)
    elif data.startswith("admin_test_atype_"):
        await set_new_test_type(bot, admin, data.replace("admin_test_atype_", ""))
    elif data.startswith("admin_test_tone_"):
        await set_new_test_tone(bot, admin, data.replace("admin_test_tone_", ""))
    elif data == "admin_new_product":
        await prompt_new_product(bot, admin)
    elif data == "admin_edit_welcome":
        await prompt_edit_welcome(bot, admin)
    elif data == "admin_edit_card":
        await prompt_edit_card(bot, admin)
    elif data == "admin_edit_card_owner":
        await prompt_edit_card_owner(bot, admin)
    elif data == "admin_export_users":
        await export_users_excel(bot, admin, db)
    elif data == "admin_static_content":
        await show_static_content_menu(bot, admin, db)
    elif data.startswith("admin_sc_list_"):
        await show_sc_list(bot, admin, data.replace("admin_sc_list_", ""), db)
    elif data.startswith("admin_sc_item_"):
        await show_sc_item(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_sc_toggle_"):
        await toggle_sc_item(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_sc_delete_"):
        await delete_sc_item(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_sc_add_"):
        await prompt_add_sc(bot, admin, data.replace("admin_sc_add_", ""))
    elif data.startswith("admin_sc_media_"):
        await sc_set_media_type(bot, admin, data.replace("admin_sc_media_", ""))
    elif data.startswith("admin_approve_pay_"):
        from bot.handlers.payment import approve_payment
        await approve_payment(bot, int(data.split("_")[-1]), db)
        await send_message(bot, admin.bale_id, "? ?????? ????? ??.")
    elif data.startswith("admin_reject_pay_"):
        await prompt_reject_payment(bot, admin, int(data.split("_")[-1]))
    elif data == "admin_discounts":
        await show_discounts(bot, admin, db)
    elif data == "admin_new_discount":
        await prompt_new_discount(bot, admin)
    elif data.startswith("admin_discount_"):
        await show_discount_item(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_toggle_discount_"):
        await toggle_discount(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_delete_discount_"):
        await delete_discount(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_view_receipt_"):
        p = db.query(Payment).filter(Payment.id == int(data.split("_")[-1])).first()
        if p and p.receipt_image:
            await send_photo(bot, admin.bale_id, p.receipt_image, f"??? {p.tracking_code}")
    elif data.startswith("admin_user_"):
        await show_user_profile(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_msg_user_"):
        await prompt_direct_message(bot, admin, int(data.split("_")[-1]))
    elif data.startswith("admin_block_"):
        u = db.query(User).filter(User.id == int(data.split("_")[-1])).first()
        if u:
            u.status = UserStatus.blocked
            db.commit()
            await send_message(bot, admin.bale_id, f"{u.full_name} ???? ??.")
    elif data.startswith("admin_unblock_"):
        u = db.query(User).filter(User.id == int(data.split("_")[-1])).first()
        if u:
            u.status = UserStatus.active
            db.commit()
            await send_message(bot, admin.bale_id, f"{u.full_name} ?????? ??.")
    elif data.startswith("admin_cancel_res_"):
        await cancel_reservation(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_reschedule_"):
        await prompt_reschedule(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_newslot_"):
        await reschedule_reservation(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_test_menu_"):
        await show_test_menu(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_toggle_test_"):
        await toggle_test(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_confirm_delete_test_"):
        await confirm_delete_test(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_delete_test_"):
        await delete_test(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_product_menu_"):
        await show_product_menu(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_toggle_product_"):
        await toggle_product(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_delete_product_"):
        await delete_product(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_ptype_"):
        await set_product_type(bot, admin, data.replace("admin_ptype_", ""))
    elif data.startswith("admin_approve_review_"):
        await approve_review(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_reject_review_"):
        await reject_review(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_approve_installment_"):
        await approve_installment_admin(bot, admin, int(data.split("_")[-1]), db)
    elif data.startswith("admin_reject_installment_"):
        await reject_installment_admin(bot, admin, int(data.split("_")[-1]))
    elif data == "admin_inst_yes":
        await handle_inst_yes(bot, admin)
    elif data == "admin_inst_no":
        await handle_inst_no(bot, admin, db)